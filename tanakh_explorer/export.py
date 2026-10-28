"""Export search results in several formats.

Every export carries a metadata section that says what was searched and how:
the kind of search, the query, every setting (text layer, maqaf, finals,
scope, skips, boundaries, method, match type …), the number of results in
total and in the file, the dates, the program version and the source text
with its license (CC BY-SA 4.0 requires attribution of derived data).

Formats: CSV, Excel (.xlsx), JSON, HTML (also printable to PDF from the
browser), Markdown and plain text.

The exporter works on the result objects the search functions return, so
anything shown in the interface can be exported, including statistics and
scan jobs that are not kept in the history.
"""

from __future__ import annotations

import csv
import html
import io
import json
import time

from . import __version__

FORMATS = {
    "csv": ("CSV", "text/csv; charset=utf-8", "csv"),
    "xlsx": ("Excel", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"),
    "json": ("JSON", "application/json; charset=utf-8", "json"),
    "html": ("HTML", "text/html; charset=utf-8", "html"),
    "md": ("Markdown", "text/markdown; charset=utf-8", "md"),
    "txt": ("טקסט", "text/plain; charset=utf-8", "txt"),
}

KIND_HE = {
    "letters": "חיפוש אותיות",
    "anagram": "חיפוש בכל סידורי האותיות (קומבינטוריקה)",
    "phrase": "חיפוש ביטוי (מילים שלמות)",
    "niqqud": "חיפוש ניקוד",
    "niqqud_sum": "חיפוש לפי גימטריית ניקוד",
    "gematria": "חיפוש גימטריה",
    "gematria_all": "חיפוש בכל סוגי הגימטריה",
    "transform": "חיפוש בחילוף אותיות",
    "transform_all": "חיפוש בכל שיטות חילופי האותיות",
    "all_combined": "חיפוש בכל הגימטריות והחילופים",
    "pipeline": "חיפוש משולב",
    "constant_scan": "סריקת כל הדילוגים הקבועים",
    "variable_scan": "סריקת כל צירופי הדילוגים המשתנים",
    "stats_letter": "סיכום ניקוד לאות",
    "stats_table": "טבלת אותיות × תנועות",
    "stats_compare": "השוואת ניקוד בין תחומים",
    "stats_summary": "סיכום תחום",
}
UNIT_HE = {"words": "מילים", "verses": "פסוקים", "letters": "אותיות"}
SELECTOR_HE = {"first": "אות ראשונה (ראשי)", "last": "אות אחרונה (סופי)", "mixed": "מעורב"}
BOUNDARY_HE = {"none": "מותר לחצות גבולות", "word": "בתוך מילה אחת", "verse": "בתוך פסוק אחד",
               "chapter": "בתוך פרק אחד", "book": "בתוך ספר אחד"}
ANCHOR_HE = {"scope": "מתחילת התחום", "verse": "מתחילת כל פסוק", "chapter": "מתחילת כל פרק"}
HL_OPEN, HL_CLOSE = "⟨", "⟩"
SOURCE_HE = ("טקסט התנ״ך: ״מקרא על פי המסורה״ (MAM), ויקיטקסט, רישיון CC BY-SA 4.0 — "
             "https://he.wikisource.org/wiki/מקרא_על_פי_המסורה . נתונים שנגזרו ממנו כפופים לאותו רישיון.")


# ----------------------------------------------------------------------------- metadata

def _fmt_list(v):
    return ", ".join(str(x) for x in v) if isinstance(v, (list, tuple)) else str(v)


def describe(kind: str, request: dict, result: dict, extra: dict | None = None) -> list[tuple[str, str]]:
    """(label, value) pairs describing the search, in Hebrew."""
    req = request or {}
    res = result or {}
    s = res.get("settings") or {}
    meta: list[tuple[str, str]] = [("סוג החיפוש", KIND_HE.get(kind, kind))]

    def add(label, value):
        if value not in (None, "", [], {}):
            meta.append((label, str(value)))

    add("כותרת", (extra or {}).get("title"))
    q = req.get("query") or s.get("query") or req.get("letter")
    add("שאילתה", _fmt_list(q) if isinstance(q, list) else q)
    if req.get("target") not in (None, ""):
        add("ערך מבוקש", req.get("target"))
    elif res.get("target") is not None:
        add("ערך מבוקש", res.get("target"))
    unit = req.get("unit") or s.get("unit")
    if unit:
        add("יחידות", UNIT_HE.get(unit, unit))
    sel = req.get("selector") or s.get("selector")
    if sel and unit in ("words", "verses"):
        add("מכל יחידה", SELECTOR_HE.get(sel, sel))
    if req.get("selectors"):
        add("בחירה לכל אות", ", ".join(SELECTOR_HE.get(x, x) for x in req["selectors"]))
    add("מצב חיפוש", s.get("search_mode_he") or res.get("search_mode_he"))
    if req.get("whole_word"):
        add("התאמה", "מילה שלמה")
    add("סוג התאמה", s.get("match_he"))
    skips = req.get("skips") or s.get("skips")
    if skips and skips != [1]:
        add("דילוג", ("קבוע " if len(skips) == 1 else "משתנה ") + _fmt_list(skips))
    elif skips == [1]:
        add("דילוג", "ללא (יחידות צמודות)")
    if req.get("skip_min") is not None:
        add("טווח דילוגים", f"{req.get('skip_min')}–{req.get('skip_max')}"
            + (" (כל צירופי הדילוגים המשתנים)" if req.get("variable") else " (דילוג קבוע)"))
    if req.get("anchor"):
        add("נקודת התחלה", f"{ANCHOR_HE.get(req['anchor'], req['anchor'])}, דילוג פתיחה {req.get('anchor_skip', 0)}")
    b = req.get("boundary") or s.get("boundary")
    if b:
        add("גבולות", BOUNDARY_HE.get(b, b))
    if req.get("reverse"):
        add("כיוון", "קדימה וגם הפוך")
    if req.get("letter_groups") == "articulation":
        add("קבוצות אותיות", "חילופי מוצאות הפה")
    method = res.get("method")
    if isinstance(method, dict):
        add("שיטת גימטריה", f"{method.get('name_he')} ({method.get('id')})")
    elif req.get("method"):
        add("שיטה", req.get("method"))
    if s.get("params"):
        add("הגדרות השיטה", json.dumps(s["params"], ensure_ascii=False))
    if req.get("span") == "verse":
        add("חיפוש ב", "פסוקים שלמים")
    elif req.get("n_min"):
        add("אורך רצף המילים", f"{req.get('n_min')}–{req.get('n_max', req.get('n_min'))}")
    tr = res.get("transformation")
    if isinstance(tr, dict):
        add("חילוף", f"{tr.get('name_he')}: {tr.get('input')} ← {tr.get('output')}")
    add("השוואה", res.get("comparison"))
    if kind == "niqqud_sum":
        add("סוג חיפוש הסכום", f"{req.get('kind', 'word')}" + (f" / {req.get('sub')}" if req.get("sub") else ""))
    if kind == "anagram":
        add("סידורים אפשריים", res.get("arrangements_total"))
        add("סידורים שנמצאו", res.get("arrangements_found"))
        add("כולל הסדר המקורי", "כן" if req.get("include_original", True) else "לא")
    lay = s.get("layer") or req.get("layer") or {}
    if s.get("layer_he"):
        add("נוסח", s["layer_he"])
    elif lay:
        add("נוסח", ("כתיב" if lay.get("kq", "ketiv") == "ketiv" else "קרי") + " · "
            + ("מקף מפריד בין מילים" if lay.get("maqaf", "split") == "split" else "צירוף במקף = מילה אחת"))
    fin = req.get("finals") or s.get("finals")
    if fin:
        add("אותיות סופיות", "התאמה מדויקת" if fin == "exact" else "ם = מ")
    add("תחום", s.get("scope_he") or res.get("scope_he"))
    add("מיקום במילה", s.get("position_he"))
    if req.get("niqqud_values"):
        add("ערכי ניקוד ששונו", json.dumps(req["niqqud_values"], ensure_ascii=False))
    # results
    total = res.get("total")
    if total is None and isinstance(res.get("rows"), list):
        total = len(res["rows"])
    total_label = "סך המופעים" if kind in ("stats_letter", "stats_table") else "סך כל התוצאות"
    add(total_label, f"{total:,}" if isinstance(total, int) else total)
    comp = res.get("complete")
    if comp is not None:
        add("שלמות", "סריקה מלאה של הטווח שהוגדר" if comp else "חלקית — הסריקה לא הושלמה")
    add("זמן החיפוש", (extra or {}).get("searched_at"))
    add("נוצר", time.strftime("%Y-%m-%d %H:%M:%S"))
    add("תוכנה", f"חוקר התנ״ך {__version__}")
    add("מקור", SOURCE_HE)
    return meta


# ----------------------------------------------------------------------------- tables

def _marked_verse(block: dict, mark=(HL_OPEN, HL_CLOSE)) -> str:
    words = []
    for w in block.get("words", []):
        t, pos, out = w["text"], 0, ""
        for s, e in sorted(w.get("hl") or []):
            if s < pos:
                continue
            out += t[pos:s] + mark[0] + t[s:e] + mark[1]
            pos = e
        out += t[pos:]
        words.append(out + ("־" if w.get("maqaf") else " "))
    return "".join(words).strip()


def _verses_text(r: dict, mark=(HL_OPEN, HL_CLOSE)) -> str:
    parts = []
    for b in r.get("display") or []:
        if "words" in b:
            parts.append(f"{b['ref']['ref_he']}: {_marked_verse(b, mark)}")
        elif b.get("more"):
            parts.append(f"(ועוד {b['more']} פסוקים)")
    return " | ".join(parts)


def _results_table(results: list[dict], title: str = "תוצאות") -> dict:
    cols = ["#", "מקום", "נמצא", "ערך", "סידור", "דילוגים", "פסוקים (האותיות שנמצאו מסומנות ⟨ ⟩)", "כיצד חושב"]
    rows = []
    for i, r in enumerate(results, 1):
        frm, to = r.get("from") or {}, r.get("to") or {}
        where = frm.get("ref_he", "")
        if to and to.get("osis") != frm.get("osis"):
            where += " – " + to.get("ref_he", "")
        rows.append([i, where, r.get("found", ""), r.get("value", ""), r.get("arrangement", ""),
                     ", ".join(map(str, r["skips"])) if r.get("skips") else "",
                     _verses_text(r), " ; ".join(r.get("steps") or [])])
    # drop columns that are empty for every row
    keep = [j for j in range(len(cols)) if j in (0, 1) or any(str(row[j]) != "" for row in rows)]
    return {"title": title, "columns": [cols[j] for j in keep], "rows": [[row[j] for j in keep] for row in rows],
            "_results": results}


def tables_for(kind: str, res: dict) -> list[dict]:
    t: list[dict] = []
    if kind == "pipeline":
        trace = res.get("trace") or []
        t.append({"title": "שלבי החיפוש", "columns": ["שלב", "פעולה", "פרט", "פלט", "חישוב"],
                  "rows": [[x.get("step"), x.get("op_he"), x.get("detail", ""), x.get("output"),
                            " ; ".join(x.get("calc") or [])] for x in trace]})
        if res.get("result"):
            t.extend(tables_for("letters", res["result"]))
        return t
    if kind == "anagram" and res.get("arrangements"):
        t.append({"title": "סידורי האותיות", "columns": ["סידור", "הופעות", "מקורי"],
                  "rows": [[a["display"], a["count"], "כן" if a.get("original") else ""] for a in res["arrangements"]]})
    if res.get("query_analysis"):
        t.append({"title": "ניתוח השאילתה", "columns": ["אות", "תנועה", "ערך ניקוד"],
                  "rows": [[q["letter"], q["vowels_he"], q["value"]] for q in res["query_analysis"]]})
    if kind == "gematria_all":
        t.append({"title": "ערכים והתאמות לפי שיטה", "columns": ["שיטה", "קטגוריה", "ערך השאילתה", "התאמות", "חישוב"],
                  "rows": [[r["name_he"], r["category"], r["value"], r["total"], " ; ".join(r.get("steps") or [])]
                           for r in res.get("rows", [])]})
    elif kind == "transform_all":
        t.append({"title": "חילופים", "columns": ["שיטה", "תוצאת החילוף", "התאמות"],
                  "rows": [[r["name_he"], r["output"], r["count"]] for r in res.get("rows", [])]})
    elif kind == "all_combined":
        t.append({"title": "חילופים × שיטות גימטריה", "columns": ["חילוף", "טקסט", "שיטה", "ערך", "מילים בעלות ערך זה"],
                  "rows": [[r["transformation_he"], r["text"], r["method_he"], r["value"], r["count"]]
                           for r in res.get("rows", [])]})
    elif kind == "constant_scan":
        t.append({"title": "התאמות לפי דילוג", "columns": ["דילוג", "כיוון", "התאמות"],
                  "rows": [[r["skip"], "הפוך" if r["direction"] == "reverse" else "קדימה", r["count"]]
                           for r in res.get("rows", [])]})
        ex = [x for r in res.get("rows", []) for x in (r.get("results") or [])]
        if ex:
            t.append(_results_table(ex, "דוגמאות"))
    elif kind == "stats_letter":
        t.append({"title": f"הניקוד על האות {res.get('letter', '')}",
                  "columns": ["תנועה", "מופעים", "אחוז", "ערך התנועה", "סכום הניקוד"],
                  "rows": [[r["name_he"], r["count"], f"{r['share'] * 100:.2f}%", r["value"] if r["value"] is not None else "",
                            r["niqqud_sum"]] for r in res.get("rows", []) if r["count"]]})
        cols = [r["column"] for r in res.get("rows", []) if r["count"]]
        names = {r["column"]: r["name_he"] for r in res.get("rows", [])}
        if res.get("by_book"):
            t.append({"title": "לפי ספר", "columns": ["ספר", "סה״כ"] + [names[c] for c in cols],
                      "rows": [[b["he"], b["total"]] + [b["cells"].get(c, 0) for c in cols] for b in res["by_book"]]})
    elif kind == "stats_table":
        cols = [c for c in res.get("columns", []) if res.get("column_totals", {}).get(c["id"])]
        t.append({"title": "אותיות × תנועות", "columns": ["אות", "סה״כ"] + [c["name_he"] for c in cols]
                  + ["סכום ניקוד", "סכום גימטריה", "משולב"],
                  "rows": [[r["letter"], r["total"]] + [r["cells"][c["id"]] for c in cols]
                           + [r["niqqud_sum"], r["gematria_sum"], r["combined_sum"]] for r in res.get("rows", []) if r["total"]]})
    elif kind == "stats_compare":
        cols = res.get("columns", [])
        t.append({"title": f"א׳ = {res.get('scope_a')} · ב׳ = {res.get('scope_b')}",
                  "columns": ["אות"] + [f"{c['name_he']} א׳/ב׳" for c in cols],
                  "rows": [[r["letter"]] + [f"{r['cells'][c['id']]['a']}/{r['cells'][c['id']]['b']}" for c in cols]
                           for r in res.get("rows", [])]})
    elif kind == "stats_summary":
        t.append({"title": "סיכום", "columns": ["נתון", "ערך"],
                  "rows": [[k, v] for k, v in res.items() if isinstance(v, (int, float, str)) and k != "scope_he"]})
        if res.get("letter_counts"):
            t.append({"title": "מספר מופעים לכל אות", "columns": ["אות", "מופעים"],
                      "rows": [[k, v] for k, v in res["letter_counts"].items()]})
    if res.get("counts_by_length"):
        t.append({"title": "התאמות לפי אורך", "columns": ["אורך", "התאמות"],
                  "rows": [[k, v] for k, v in res["counts_by_length"].items()]})
    if isinstance(res.get("results"), list) and kind != "constant_scan":
        t.append(_results_table(res["results"]))
    return t


# ----------------------------------------------------------------------------- writers

def _csv(meta, tables) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["מידע על החיפוש"])
    for k, v in meta:
        w.writerow([k, v])
    for t in tables:
        w.writerow([])
        w.writerow([t["title"]])
        w.writerow(t["columns"])
        w.writerows(t["rows"])
    return ("\ufeff" + buf.getvalue()).encode("utf-8")


def _xlsx(meta, tables) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "מידע על החיפוש"
    ws.sheet_view.rightToLeft = True
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="E6EEF6")
    for k, v in meta:
        ws.append([k, v])
        ws.cell(ws.max_row, 1).font = bold
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 90
    for row in ws.iter_rows():
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    used = set()
    for t in tables:
        name = "".join(ch for ch in t["title"] if ch not in "[]:*?/\\")[:28] or "טבלה"
        base, i = name, 2
        while name in used:
            name = f"{base[:25]} {i}"
            i += 1
        used.add(name)
        sh = wb.create_sheet(name)
        sh.sheet_view.rightToLeft = True
        sh.append(t["columns"])
        for c in sh[1]:
            c.font = bold
            c.fill = head_fill
        for r in t["rows"]:
            sh.append([x if isinstance(x, (int, float)) else ("" if x is None else str(x)) for x in r])
        for j, col in enumerate(t["columns"], 1):
            width = max([len(str(col))] + [min(len(str(r[j - 1])), 80) for r in t["rows"][:200]]) + 2
            sh.column_dimensions[sh.cell(1, j).column_letter].width = min(max(width, 6), 80)
        sh.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _json(meta, tables, kind, request, result) -> bytes:
    clean = [{k: v for k, v in t.items() if not k.startswith("_")} for t in tables]
    doc = {"metadata": dict(meta), "kind": kind, "request": request, "tables": clean, "result": result}
    return json.dumps(doc, ensure_ascii=False, indent=1, default=str).encode("utf-8")


def _html(meta, tables) -> bytes:
    e = html.escape
    title = next((v for k, v in meta if k == "כותרת"), None) or next(v for k, v in meta if k == "סוג החיפוש")
    parts = [f"""<!doctype html><html lang="he" dir="rtl"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1"><title>{e(str(title))}</title>
<style>
body{{font-family:"Noto Serif Hebrew","Taamey Frank CLM","SBL Hebrew","David","Times New Roman",serif;margin:24px;color:#1d2433;background:#fff}}
h1{{font-size:22px;margin:0 0 12px}} h2{{font-size:17px;margin:22px 0 8px;color:#2f5d8a}}
table{{border-collapse:collapse;width:100%;font-size:14px;margin-bottom:8px}}
th,td{{border:1px solid #d9d4c8;padding:5px 7px;vertical-align:top;text-align:right}}
th{{background:#e6eef6}} table.meta th{{width:180px}}
td.v{{font-size:18px;line-height:1.9}} mark{{background:#ffd86b;border-radius:3px}}
.note{{color:#5f6b7d;font-size:13px}}
@media print{{body{{margin:10mm}} h2{{break-after:avoid}} tr{{break-inside:avoid}}}}
</style></head><body><h1>{e(str(title))}</h1>
<p class="note">להפקת PDF: הדפיסו את הדף מהדפדפן ובחרו ״שמירה כ־PDF״.</p>
<table class="meta">"""]
    for k, v in meta:
        parts.append(f"<tr><th>{e(k)}</th><td>{e(str(v))}</td></tr>")
    parts.append("</table>")
    for t in tables:
        parts.append(f"<h2>{e(t['title'])}</h2><table><tr>" + "".join(f"<th>{e(str(c))}</th>" for c in t["columns"]) + "</tr>")
        results = t.get("_results")
        for i, r in enumerate(t["rows"]):
            cells = []
            for c, x in zip(t["columns"], r):
                if results is not None and c.startswith("פסוקים"):
                    blocks = results[i].get("display") or []
                    inner = "<br>".join(
                        f"<span class='note'>{e(b['ref']['ref_he'])}</span> " + _html_verse(b) if "words" in b
                        else f"<span class='note'>(ועוד {b.get('more')} פסוקים)</span>" for b in blocks)
                    cells.append(f"<td class='v'>{inner}</td>")
                else:
                    cells.append(f"<td>{e('' if x is None else str(x))}</td>")
            parts.append("<tr>" + "".join(cells) + "</tr>")
        parts.append("</table>")
    parts.append("</body></html>")
    return "".join(parts).encode("utf-8")


def _html_verse(block: dict) -> str:
    out = []
    for w in block.get("words", []):
        t, pos, s_out = w["text"], 0, ""
        for s, e_ in sorted(w.get("hl") or []):
            if s < pos:
                continue
            s_out += html.escape(t[pos:s]) + "<mark>" + html.escape(t[s:e_]) + "</mark>"
            pos = e_
        s_out += html.escape(t[pos:])
        out.append(s_out + ("־" if w.get("maqaf") else " "))
    return "".join(out)


def _md_cell(x) -> str:
    return str("" if x is None else x).replace("|", "\\|").replace("\n", " ")


def _md(meta, tables) -> bytes:
    title = next(v for k, v in meta if k == "סוג החיפוש")
    lines = [f"# {title}", "", "| שדה | ערך |", "|---|---|"]
    lines += [f"| {_md_cell(k)} | {_md_cell(v)} |" for k, v in meta]
    for t in tables:
        lines += ["", f"## {t['title']}", "", "| " + " | ".join(map(_md_cell, t["columns"])) + " |",
                  "|" + "---|" * len(t["columns"])]
        lines += ["| " + " | ".join(_md_cell(x) for x in r) + " |" for r in t["rows"]]
    return ("\n".join(lines) + "\n").encode("utf-8")


def _txt(meta, tables) -> bytes:
    lines = ["מידע על החיפוש", "=" * 30]
    width = max(len(k) for k, _ in meta)
    lines += [f"{k.ljust(width)} : {v}" for k, v in meta]
    for t in tables:
        lines += ["", t["title"], "-" * 30]
        for r in t["rows"]:
            lines.append(" · ".join(f"{c}: {x}" for c, x in zip(t["columns"], r) if x not in ("", None)))
    return ("\ufeff" + "\n".join(lines) + "\n").encode("utf-8")


def export(fmt: str, kind: str, request: dict, result: dict, extra: dict | None = None) -> tuple[bytes, str, str]:
    """Return (content, media type, file extension)."""
    if fmt not in FORMATS:
        raise ValueError(f"פורמט לא נתמך: {fmt}")
    meta = describe(kind, request, result, extra)
    tables = tables_for(kind, result)
    exported = next((len(t["rows"]) for t in tables if "_results" in t), None)
    if exported is not None:
        meta.insert(next(i for i, (k, _) in enumerate(meta) if k == "סך כל התוצאות") + 1
                    if any(k == "סך כל התוצאות" for k, _ in meta) else len(meta),
                    ("תוצאות בקובץ זה", f"{exported:,}"))
    total = result.get("total") if isinstance(result, dict) else None
    if exported is not None and isinstance(total, int) and exported < total:
        meta.insert(1, ("הערה", f"הקובץ כולל {exported:,} מתוך {total:,} התוצאות. לייצוא כל התוצאות "
                               "(עד 5,000) סמנו ״כל התוצאות״ בחלונית הייצוא."))
    if fmt == "csv":
        data = _csv(meta, tables)
    elif fmt == "xlsx":
        data = _xlsx(meta, tables)
    elif fmt == "json":
        data = _json(meta, tables, kind, request, result)
    elif fmt == "html":
        data = _html(meta, tables)
    elif fmt == "md":
        data = _md(meta, tables)
    else:
        data = _txt(meta, tables)
    _, media, ext = FORMATS[fmt]
    return data, media, ext
