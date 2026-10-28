import pytest

from tanakh_explorer.combined import pipeline


def test_pipeline_type_checking():
    assert pipeline.validate([{"op": "transform", "transformation": "atbash"}, {"op": "gematria"},
                              {"op": "search_value"}]) == []
    errs = pipeline.validate([{"op": "gematria"}, {"op": "transform", "transformation": "atbash"}])
    assert errs and "מצפה לטקסט" in errs[0]
    assert pipeline.validate([{"op": "search_value"}])  # number expected, text given
    assert pipeline.validate([{"op": "nope"}])


def test_pipeline_runs(synth):
    res = pipeline.run(synth, {"query": "שָׁלוֹם", "steps": [
        {"op": "transform", "transformation": "atbash"},
        {"op": "gematria", "method": "kidmi", "include_niqqud": True},
        {"op": "arith", "operator": "+", "value": 1},
        {"op": "search_value", "method": "standard"},
    ]})
    assert res["ok"]
    ops = [t["op"] for t in res["trace"]]
    assert ops == ["input", "transform", "gematria", "arith", "search_value"]
    # kidmi(בכפי)=3+75+405+55=538, niqqud 26, +1
    assert res["trace"][2]["output"] == 564 and res["trace"][3]["output"] == 565


@pytest.fixture(scope="module")
def client(corpus):
    from fastapi.testclient import TestClient
    from tanakh_explorer.api.app import app
    return TestClient(app)


def test_api_endpoints(client):
    m = client.get("/api/meta").json()
    assert len(m["books"]) == 39 and m["gematria_methods"] and m["transformations"]
    r = client.post("/api/search/letters", json={"query": "שלום", "unit": "words", "limit": 2, "save": False}).json()
    assert r["total"] > 0 and "history_id" not in r
    r = client.post("/api/search/letters", json={"query": "", "save": False})
    assert r.status_code == 400 and "error" in r.json()
    r = client.post("/api/calc/gematria", json={"query": "שלום", "methods": ["standard", "gadol"]}).json()
    assert [x["value"] for x in r] == [376, 936]
    r = client.get("/api/reader/Gen/1").json()
    assert len(r["verses"]) == 31
    r = client.post("/api/search/gematria", json={"target": 376, "limit": 1}).json()
    hid = r["history_id"]
    csv = client.get(f"/api/history/{hid}/export?format=csv")
    assert csv.status_code == 200 and "text/csv" in csv.headers["content-type"]
    client.delete(f"/api/history/{hid}")


def test_export_all_reruns_search(client):
    import io
    from openpyxl import load_workbook
    r = client.post("/api/search/letters", json={"query": "שלום", "unit": "words", "limit": 2}).json()
    assert r["total"] > 2
    full = client.post("/api/export", json={"format": "xlsx", "kind": "letters", "request": {"query": "שלום", "unit": "words"},
                                            "result": r, "all": True})
    assert full.status_code == 200 and "filename*=UTF-8''" in full.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(full.content))
    assert wb["תוצאות"].max_row == r["total"] + 1  # header + every result
    part = client.get(f"/api/history/{r['history_id']}/export?format=md")
    assert part.status_code == 200 and "סוג החיפוש" in part.content.decode()
    client.delete(f"/api/history/{r['history_id']}")
