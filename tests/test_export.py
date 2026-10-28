import io
import json

import pytest

from tanakh_explorer import export as E
from tanakh_explorer.search import service as S
from tanakh_explorer.stats import analysis as A


def _search_results(synth):
    return {
        "letters": ({"query": "אב", "unit": "letters", "limit": 5}, S.letter_search(synth, {"query": "אב", "unit": "letters", "limit": 5})),
        "gematria": ({"target": 3}, S.gematria_search(synth, {"target": 3, "limit": 5})),
        "anagram": ({"query": "אבג", "search_mode": "sequence"}, S.anagram_search(synth, {"query": "אבג", "search_mode": "sequence", "limit": 5})),
        "gematria_all": ({"query": "אב"}, S.gematria_all(synth, {"query": "אב", "methods": ["standard", "ordinal"]})),
        "stats_letter": ({"letter": "א"}, A.letter_summary(synth, {"letter": "א", "layer": {"kq": "qere"}})),
    }


@pytest.mark.parametrize("fmt", list(E.FORMATS))
def test_every_format_carries_metadata(synth, fmt):
    for kind, (req, res) in _search_results(synth).items():
        data, media, ext = E.export(fmt, kind, req, res, {"title": "בדיקה"})
        assert data and ext == E.FORMATS[fmt][2]
        if fmt == "xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(data))
            info = wb["מידע על החיפוש"]
            assert info.sheet_view.rightToLeft
            labels = [r[0].value for r in info.iter_rows()]
        else:
            text = data.decode("utf-8-sig")
            if fmt == "json":
                labels = list(json.loads(text)["metadata"])
            else:
                labels = [k for k, _ in E.describe(kind, req, res)]
                for k in ("סוג החיפוש", "מקור"):
                    assert k in text
        assert "סוג החיפוש" in labels and "מקור" in labels
        assert any(lb in labels for lb in ("סך כל התוצאות", "סך המופעים"))


def test_partial_export_is_labelled(synth):
    res = S.letter_search(synth, {"query": "א", "unit": "letters", "limit": 2})
    assert res["total"] > 2
    data, _, _ = E.export("txt", "letters", {"query": "א"}, res)
    text = data.decode("utf-8-sig")
    assert f"הקובץ כולל 2 מתוך {res['total']:,}" in text


def test_highlight_marks_found_letters(synth):
    res = S.letter_search(synth, {"query": "אב", "unit": "letters", "limit": 1})
    table = E.tables_for("letters", res)[-1]
    verses = table["rows"][0][table["columns"].index(next(c for c in table["columns"] if c.startswith("פסוקים")))]
    assert E.HL_OPEN in verses and E.HL_CLOSE in verses
